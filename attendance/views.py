# attendance/views.py

from datetime import datetime, timedelta, date
from calendar import monthrange
from io import BytesIO
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict, OrderedDict
from urllib.parse import urlencode, quote as urlquote
import os
import uuid
from django.db.models import Sum
import qrcode

from .models import ProjectLunchRule
from .forms import ProjectLunchRuleForm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.drawing.image import Image as XLImage

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.core.paginator import Paginator
from django.core.signing import TimestampSigner, SignatureExpired, BadSignature
from django.db import transaction, IntegrityError
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.csrf import csrf_protect

from .forms import (
    EmployeeForm,
    ScanForm,
    PayrollPeriodForm,
    PayrollLineAdjustmentForm,
    PayrollExportForm,
    ProjectForm,
    IncomeForm,
    ExpenseForm,
    ProjectLunchRuleForm,
)

from .models import (
    Employee,
    Attendance,
    AttendanceScan,
    Project,
    PayrollPeriod,
    PayrollLine,
    ProjectIncome,
    ProjectExpense,
    ProjectLunchRule,
)


# ──────────────────────────────────────────────────────────────────────────────
# Helpers generales
# ──────────────────────────────────────────────────────────────────────────────

def is_admin(user):
    return user.is_superuser


def _user_can_access_project(user, project: Project) -> bool:
    return user.is_superuser or project.supervisors.filter(id=user.id).exists()


def _available_projects(user):
    if user.is_superuser:
        return Project.objects.all().order_by("name")
    return user.supervised_projects.all().order_by("name")


def _to_decimal(value, default="0"):
    if value is None:
        return Decimal(default)
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _money_prefix(currency):
    return currency or "USD"


def format_money_latam(value, currency="USD") -> str:
    q = Decimal(value or 0).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    us = f"{q:,.2f}"
    latam = us.replace(",", "§").replace(".", ",").replace("§", ".")
    return f"{_money_prefix(currency)} {latam}"


def format_money_latam_whole(value, currency="USD") -> str:
    q = Decimal(value or 0).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    us = f"{q:,.0f}"
    latam = us.replace(",", "§").replace(".", ",").replace("§", ".")
    return f"{_money_prefix(currency)} {latam}"


def format_hhmm_from_minutes(minutes) -> str:
    m = int(Decimal(minutes or 0))
    return f"{m // 60:02d}:{m % 60:02d}"


def minutes_to_hours_decimal(minutes: int) -> Decimal:
    return (Decimal(minutes or 0) / Decimal(60)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _parse_hhmm_to_minutes(value: str) -> int:
    value = (value or "").strip()

    if not value:
        raise ValueError("Vacío")

    if ":" in value:
        hh, mm = value.split(":", 1)
        h = int(hh)
        m = int(mm)

        if h < 0 or not (0 <= m < 60):
            raise ValueError("Hora inválida")

        return h * 60 + m

    q = Decimal(value)
    if q < 0:
        raise ValueError("Valor negativo")

    return int((q * Decimal(60)).to_integral_value(rounding=ROUND_HALF_UP))


def _billable_minutes(att: Attendance) -> int:
    try:
        return int(att.worked_minutes())
    except Exception:
        return 0


def _parse_device_ts(device_ts, fallback):
    if not device_ts:
        return fallback

    try:
        dt = datetime.fromisoformat(device_ts.replace("Z", "+00:00"))

        if timezone.is_naive(dt):
            dt = timezone.make_aware(dt)

        return dt
    except Exception:
        return fallback


def _wants_json(request):
    accept = request.headers.get("Accept", "")
    content_type = request.headers.get("Content-Type", "")
    return "application/json" in accept or "application/json" in content_type


def _add_logo(ws, cell="A1", max_width_px=150):
    logo_path = os.path.join(settings.BASE_DIR, "static", "img", "LOGO_PERFIL_AMBU_GROUP.png")

    if not os.path.exists(logo_path):
        return

    ws.row_dimensions[1].height = 85
    ws.column_dimensions["A"].width = 23

    img = XLImage(logo_path)

    try:
        from PIL import Image as PILImage

        with PILImage.open(logo_path) as im:
            w, h = im.size

        if w > max_width_px:
            ratio = max_width_px / float(w)
            img.width = int(w * ratio)
            img.height = int(h * ratio)
    except Exception:
        pass

    ws.add_image(img, cell)


def _month_range_from_request(request):
    today = timezone.localdate()
    raw = (request.GET.get("month") or "").strip()

    try:
        if raw:
            y, m = [int(x) for x in raw.split("-")]
        else:
            y, m = today.year, today.month
    except Exception:
        y, m = today.year, today.month

    first = date(y, m, 1)
    last = date(y, m, monthrange(y, m)[1])

    return y, m, first, last


def _paginate(request, queryset, per_page=25):
    paginator = Paginator(queryset, per_page)
    page_number = request.GET.get("page") or 1
    return paginator.get_page(page_number)


# ──────────────────────────────────────────────────────────────────────────────
# Dashboard
# ──────────────────────────────────────────────────────────────────────────────

@login_required
def dashboard_view(request):
    projects = _available_projects(request.user)
    return render(request, "attendance/dashboard.html", {"projects": projects})


# ──────────────────────────────────────────────────────────────────────────────
# Scan QR / Biométrico
# ──────────────────────────────────────────────────────────────────────────────

def _register_attendance_scan(request, project, form, source_default="qr"):
    employee = form.cleaned_data["employee"]
    action = form.cleaned_data["action"]
    photo = form.cleaned_data["evidence_photo"]

    source = form.cleaned_data.get("source") or source_default
    verification_reference = form.cleaned_data.get("verification_reference") or ""
    device_name = form.cleaned_data.get("device_name") or ""

    client_uuid = form.cleaned_data.get("client_uuid") or str(uuid.uuid4())
    device_ts_raw = form.cleaned_data.get("device_ts")

    now = timezone.now()
    today = timezone.localdate()
    device_ts = _parse_device_ts(device_ts_raw, now)

    project.employees.add(employee)

    try:
        AttendanceScan.objects.create(
            client_uuid=client_uuid,
            employee=employee,
            project=project,
            action=("in" if action == "IN" else "out"),
            device_ts=device_ts,
            source=source,
            device_name=device_name[:100],
            verification_reference=verification_reference[:120],
        )
        duplicate = False
    except IntegrityError:
        duplicate = True

    att, _ = Attendance.objects.select_for_update().get_or_create(
        project=project,
        employee=employee,
        date=today,
    )

    att.evidence_photo = photo

    if action == "IN":
        if att.check_in is None:
            att.check_in = now
            att.check_in_source = source
            att.save()
            msg = f"Ingreso registrado: {employee.full_name} a las {timezone.localtime(now).strftime('%H:%M')}."
        else:
            msg = f"{employee.full_name} ya tenía ingreso registrado a las {timezone.localtime(att.check_in).strftime('%H:%M')}."

    else:
        if att.check_in is None:
            return {
                "ok": False,
                "status": 409,
                "duplicate": duplicate,
                "message": "No hay ingreso previo. Primero registra la entrada.",
                "redirect_error": True,
            }

        if att.check_out is None:
            att.check_out = now
            att.check_out_source = source
            att.save()

        net_minutes = _billable_minutes(att)
        msg = (
            f"Salida registrada: {employee.full_name} a las {timezone.localtime(now).strftime('%H:%M')}. "
            f"Tiempo neto trabajado: {format_hhmm_from_minutes(net_minutes)}."
        )

    return {
        "ok": True,
        "status": 200,
        "duplicate": duplicate,
        "message": msg,
    }


@login_required
@transaction.atomic
def scan_view(request, project_slug):
    project = get_object_or_404(Project, slug=project_slug, active=True)

    if not _user_can_access_project(request.user, project):
        raise Http404("Proyecto no disponible para tu usuario.")

    if request.method == "POST":
        form = ScanForm(request.POST, request.FILES, project=project)

        if not form.is_valid():
            if _wants_json(request):
                return JsonResponse({
                    "ok": False,
                    "error": "Formulario inválido.",
                    "details": form.errors,
                }, status=400)

            messages.error(request, "Revisa los datos. Hay campos incompletos o inválidos.")
            return redirect("scan", project_slug=project.slug)

        result = _register_attendance_scan(request, project, form, source_default="qr")

        if _wants_json(request):
            return JsonResponse({
                "ok": result["ok"],
                "duplicate": result.get("duplicate", False),
                "message": result["message"],
            }, status=result["status"])

        if result["ok"]:
            messages.success(request, result["message"])
        else:
            messages.error(request, result["message"])

        return redirect("scan", project_slug=project.slug)

    form = ScanForm(project=project)
    return render(request, "attendance/scan.html", {
        "form": form,
        "project": project,
    })


@login_required
def qr_scan_view(request, project_slug):
    project = get_object_or_404(Project, slug=project_slug, active=True)

    if not _user_can_access_project(request.user, project):
        raise Http404()

    signer = TimestampSigner(salt="kiosk-scan-v1")
    today_str = timezone.localdate().isoformat()
    payload = f"{project.id}:{project.qr_secret}:{today_str}"
    token = signer.sign(payload)

    kiosk_url = request.build_absolute_uri(
        reverse("scan_kiosk", args=[project.slug]) + "?" + urlencode({"t": token})
    )

    qr = qrcode.QRCode(version=1, box_size=10, border=4)
    qr.add_data(kiosk_url)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")

    stream = BytesIO()
    img.save(stream, format="PNG")

    return HttpResponse(stream.getvalue(), content_type="image/png")


@csrf_protect
@transaction.atomic
def scan_kiosk_view(request, project_slug):
    project = get_object_or_404(Project, slug=project_slug, active=True)

    token = request.GET.get("t")

    if not token:
        return HttpResponse("Token requerido.", status=403)

    signer = TimestampSigner(salt="kiosk-scan-v1")

    try:
        data = signer.unsign(token, max_age=36 * 3600)
        parts = data.split(":")

        if len(parts) != 3:
            return HttpResponse("Token inválido.", status=403)

        proj_id, secret, token_date = parts
        today_str = timezone.localdate().isoformat()

        if str(project.id) != proj_id or secret != project.qr_secret:
            return HttpResponse("Token inválido.", status=403)

        if token_date != today_str:
            return HttpResponse("Token expirado para el día de hoy.", status=403)

    except SignatureExpired:
        return HttpResponse("Token expirado.", status=403)

    except BadSignature:
        return HttpResponse("Token inválido.", status=403)

    if request.method == "POST":
        form = ScanForm(request.POST, request.FILES, project=project)

        if not form.is_valid():
            if _wants_json(request):
                return JsonResponse({
                    "ok": False,
                    "error": "Formulario inválido.",
                    "details": form.errors,
                }, status=400)

            messages.error(request, "Revisa los datos. Hay campos incompletos o inválidos.")
            return redirect(f"{reverse('scan_kiosk', args=[project.slug])}?t={urlquote(token)}")

        result = _register_attendance_scan(request, project, form, source_default="qr")

        if _wants_json(request):
            return JsonResponse({
                "ok": result["ok"],
                "duplicate": result.get("duplicate", False),
                "message": result["message"],
            }, status=result["status"])

        if result["ok"]:
            messages.success(request, result["message"])
        else:
            messages.error(request, result["message"])

        return redirect(f"{reverse('scan_kiosk', args=[project.slug])}?t={urlquote(token)}")

    form = ScanForm(project=project)

    return render(request, "attendance/scan_kiosk.html", {
        "form": form,
        "project": project,
        "kiosk": True,
    })


# ──────────────────────────────────────────────────────────────────────────────
# Reporte de asistencia
# ──────────────────────────────────────────────────────────────────────────────

@login_required
def report_view(request):
    if request.method == "POST" and request.user.is_superuser:
        att_id = request.POST.get("att_id")
        action = request.POST.get("action")

        att = get_object_or_404(
            Attendance.objects.select_related("employee", "project"),
            pk=att_id,
        )

        if action == "save":
            try:
                minutes = _parse_hhmm_to_minutes(request.POST.get("hhmm"))
            except Exception:
                messages.error(request, "Formato inválido. Usa HH:MM, ejemplo 08:00, o decimal, ejemplo 8.5.")
                return redirect("report")

            note = (request.POST.get("note") or "").strip()

            att.manual_minutes = minutes
            att.manual_note = note[:200]
            att.save(update_fields=["manual_minutes", "manual_note", "updated_at"])

            messages.success(
                request,
                f"Horas ajustadas a {format_hhmm_from_minutes(minutes)} para {att.employee.full_name}."
            )

        elif action == "clear":
            att.manual_minutes = None
            att.manual_note = ""
            att.save(update_fields=["manual_minutes", "manual_note", "updated_at"])
            messages.success(request, f"Se eliminó el ajuste manual de {att.employee.full_name}.")

        else:
            messages.error(request, "Acción no válida.")

        return redirect("report")

    q_name = (request.GET.get("name") or "").strip()
    q_document = (request.GET.get("document") or "").strip()
    q_proj = (request.GET.get("project") or "").strip()
    q_source = (request.GET.get("source") or "").strip()

    y, m, first, last = _month_range_from_request(request)

    qs = Attendance.objects.select_related("employee", "project").filter(
        date__range=(first, last)
    )

    if not request.user.is_superuser:
        qs = qs.filter(project__in=_available_projects(request.user))

    if q_proj:
        qs = qs.filter(project__slug=q_proj)

    if q_name:
        qs = qs.filter(employee__full_name__icontains=q_name)

    if q_document:
        qs = qs.filter(employee__document_id__icontains=q_document)

    if q_source:
        qs = qs.filter(check_in_source=q_source) | qs.filter(check_out_source=q_source)

    qs = qs.order_by("-date", "project__name", "employee__full_name")

    for a in qs:
        if a.check_in and a.check_out:
            gross = int((a.check_out - a.check_in).total_seconds() // 60)
        else:
            gross = 0

        a.gross_hhmm = format_hhmm_from_minutes(gross)
        a.worked_minutes_net = _billable_minutes(a)
        a.worked_hhmm_net = format_hhmm_from_minutes(a.worked_minutes_net)

    page_obj = _paginate(request, qs, per_page=25)

    totals = {}
    grand_minutes = 0
    grand_pay_by_currency = defaultdict(lambda: Decimal("0.00"))

    for a in page_obj.object_list:
        currency = a.employee.currency or "USD"
        rate = a.employee.hourly_rate or Decimal("0.00")
        minutes = _billable_minutes(a)
        pay = (Decimal(minutes) / Decimal(60)) * rate

        key = (a.project.slug, a.employee.id)

        data = totals.setdefault(key, {
            "project": a.project.name,
            "name": a.employee.full_name,
            "document": a.employee.document_id,
            "minutes": 0,
            "rate": rate,
            "currency": currency,
            "pay": Decimal("0.00"),
        })

        data["minutes"] += minutes
        data["pay"] += pay

    for info in totals.values():
        info["hhmm"] = format_hhmm_from_minutes(info["minutes"])
        info["rate_fmt"] = format_money_latam(info["rate"], info["currency"])
        info["pay_fmt"] = format_money_latam(info["pay"], info["currency"])
        grand_minutes += info["minutes"]
        grand_pay_by_currency[info["currency"]] += info["pay"]

    grand_pay_fmt = [
        format_money_latam(value, currency)
        for currency, value in grand_pay_by_currency.items()
    ]

    if not request.user.is_superuser:
        for info in totals.values():
            info["rate"] = None
            info["pay"] = None
            info["rate_fmt"] = None
            info["pay_fmt"] = None
        grand_pay_fmt = []

    return render(request, "attendance/report.html", {
        "rows": page_obj.object_list,
        "page_obj": page_obj,
        "totals": totals,
        "grand_hhmm": format_hhmm_from_minutes(grand_minutes),
        "grand_pay_fmt": grand_pay_fmt,
        "month": f"{y}-{m:02d}",
        "q_name": q_name,
        "q_document": q_document,
        "q_proj": q_proj,
        "q_source": q_source,
        "projects": _available_projects(request.user),
    })


# ──────────────────────────────────────────────────────────────────────────────
# Empleados
# ──────────────────────────────────────────────────────────────────────────────

@login_required
def employee_list(request):
    projects_qs = _available_projects(request.user)

    if request.user.is_superuser:
        employees = Employee.objects.all()
    else:
        employees = Employee.objects.filter(projects__in=projects_qs).distinct()

    q_proj = (request.GET.get("project") or "").strip()
    q_name = (request.GET.get("name") or "").strip()
    q_document = (request.GET.get("document") or "").strip()
    q_active = (request.GET.get("active") or "").strip()

    if q_proj:
        employees = employees.filter(projects__slug=q_proj)

    if q_name:
        employees = employees.filter(full_name__icontains=q_name)

    if q_document:
        employees = employees.filter(document_id__icontains=q_document)

    if q_active == "1":
        employees = employees.filter(active=True)
    elif q_active == "0":
        employees = employees.filter(active=False)

    employees = employees.order_by("full_name").prefetch_related("projects")

    page_obj = _paginate(request, employees, per_page=25)

    return render(request, "attendance/employee_list.html", {
        "employees": page_obj.object_list,
        "page_obj": page_obj,
        "projects": projects_qs,
        "q_proj": q_proj,
        "q_name": q_name,
        "q_document": q_document,
        "q_active": q_active,
    })


@login_required
def employee_form(request, pk=None):
    emp = get_object_or_404(Employee, pk=pk) if pk else None

    if request.method == "POST":
        form = EmployeeForm(request.POST, instance=emp)

        if form.is_valid():
            form.save()
            messages.success(request, "Empleado guardado correctamente.")
            return redirect("employee_list")

        messages.error(request, "No se pudo guardar. Revisa los campos marcados.")

    else:
        form = EmployeeForm(instance=emp)

    return render(request, "attendance/employee_form.html", {
        "form": form,
        "emp": emp,
    })


# ──────────────────────────────────────────────────────────────────────────────
# Reglas de almuerzo
# ──────────────────────────────────────────────────────────────────────────────

@login_required
@user_passes_test(is_admin)
def lunch_rule_list(request):
    q_proj = (request.GET.get("project") or "").strip()

    rules = ProjectLunchRule.objects.select_related("project").order_by("project__name", "weekday")

    if q_proj:
        rules = rules.filter(project__slug=q_proj)

    page_obj = _paginate(request, rules, per_page=25)

    return render(request, "attendance/lunch_rule_list.html", {
        "rules": page_obj.object_list,
        "page_obj": page_obj,
        "projects": Project.objects.all().order_by("name"),
        "q_proj": q_proj,
    })


@login_required
@user_passes_test(is_admin)
def lunch_rule_form(request, pk=None):
    rule = get_object_or_404(ProjectLunchRule, pk=pk) if pk else None

    if request.method == "POST":
        form = ProjectLunchRuleForm(request.POST, instance=rule)

        if form.is_valid():
            form.save()
            messages.success(request, "Regla de almuerzo guardada correctamente.")
            return redirect("lunch_rule_list")

        messages.error(request, "No se pudo guardar la regla. Revisa los campos.")

    else:
        form = ProjectLunchRuleForm(instance=rule)

    return render(request, "attendance/lunch_rule_form.html", {
        "form": form,
        "rule": rule,
    })


# ──────────────────────────────────────────────────────────────────────────────
# Nómina
# ──────────────────────────────────────────────────────────────────────────────

TWOPL = Decimal("0.01")


def _monday(d: date) -> date:
    return d - timedelta(days=d.weekday())


def _sunday(d: date) -> date:
    return d + timedelta(days=(6 - d.weekday()))


def _week_chunks(start: date, end: date):
    cur = _monday(start)
    last = _sunday(end)
    chunks = []

    while cur <= last:
        w_start = cur
        w_end = cur + timedelta(days=6)
        chunks.append((w_start, w_end))
        cur = w_end + timedelta(days=1)

    return chunks


def _build_week_table(employees, att_map, week_start, week_end, period_start, period_end):
    cols = [week_start + timedelta(days=i) for i in range(7)]
    col_labels = [d.strftime("%a %d") for d in cols]

    rows = []
    totals_by_currency = defaultdict(lambda: Decimal("0.00"))
    week_hours_all = Decimal("0.00")

    for emp in employees:
        currency = emp.currency or "USD"
        rate = _to_decimal(emp.hourly_rate, "0")

        day_hours_hhmm = []
        total_minutes = 0

        for d in cols:
            if period_start <= d <= period_end:
                mins = int(att_map.get((emp.id, d), 0))
            else:
                mins = 0

            day_hours_hhmm.append(format_hhmm_from_minutes(mins))
            total_minutes += mins

        total_hours = Decimal(total_minutes) / Decimal(60)
        total_pay = (total_hours * rate).quantize(TWOPL, rounding=ROUND_HALF_UP)

        rows.append({
            "name": emp.full_name,
            "document": emp.document_id,
            "day_hours_hhmm": day_hours_hhmm,
            "total_hours_hhmm": format_hhmm_from_minutes(total_minutes),
            "hourly": format_money_latam(rate, currency),
            "total_pay": format_money_latam_whole(total_pay, currency),
            "currency": currency,
            "total_pay_raw": total_pay,
        })

        week_hours_all += total_hours
        totals_by_currency[currency] += total_pay

    return {
        "columns": col_labels,
        "rows": rows,
        "total_hours_all_hhmm": format_hhmm_from_minutes(int(week_hours_all * Decimal(60))),
        "total_payroll": [format_money_latam_whole(v, c) for c, v in totals_by_currency.items()],
        "total_hours_all_raw": week_hours_all,
        "total_payroll_raw": totals_by_currency,
        "week_start": week_start,
        "week_end": week_end,
    }


@login_required
@user_passes_test(is_admin)
def payroll_list(request):
    q_project = (request.GET.get("project") or "").strip()
    q_status = (request.GET.get("status") or "").strip()

    qs = PayrollPeriod.objects.select_related("project").order_by("-start_date", "-id")

    if q_project:
        qs = qs.filter(project__slug=q_project)

    if q_status in ["OPEN", "CLOSED"]:
        qs = qs.filter(status=q_status)

    page_obj = _paginate(request, qs, per_page=20)

    return render(request, "attendance/payroll_list.html", {
        "periods": page_obj.object_list,
        "page_obj": page_obj,
        "projects": Project.objects.all().order_by("name"),
        "q_project": q_project,
        "q_status": q_status,
    })


@login_required
@user_passes_test(is_admin)
def payroll_new(request):
    today = timezone.localdate()

    if request.method == "POST":
        form = PayrollPeriodForm(request.POST)

        if form.is_valid():
            period = form.save(commit=False)
            period.created_by = request.user
            period.save()

            messages.success(request, "Periodo creado. Ahora se calculará la nómina.")
            return redirect("payroll_recalc", pk=period.pk)

        messages.error(request, "No se pudo crear el periodo. Revisa las fechas.")

    else:
        y, m = today.year, today.month
        last_day = monthrange(y, m)[1]

        if today.day <= 15:
            initial = {"start_date": date(y, m, 1), "end_date": date(y, m, 15)}
        else:
            initial = {"start_date": date(y, m, 16), "end_date": date(y, m, last_day)}

        form = PayrollPeriodForm(initial=initial)

    return render(request, "attendance/payroll_new.html", {"form": form})


@login_required
@user_passes_test(is_admin)
def payroll_recalc(request, pk):
    period = get_object_or_404(PayrollPeriod, pk=pk)

    if not period.is_open:
        messages.warning(request, "Este periodo está cerrado. No se puede recalcular.")
        return redirect("payroll_detail", pk=period.pk)

    qs = Attendance.objects.select_related("employee", "project").filter(
        date__range=(period.start_date, period.end_date)
    )

    if period.project_id:
        qs = qs.filter(project=period.project)

    minutes_map = defaultdict(int)

    for att in qs:
        minutes_map[att.employee_id] += _billable_minutes(att)

    existing = {line.employee_id: line for line in period.lines.all()}
    period.lines.exclude(employee_id__in=minutes_map.keys()).delete()

    for emp_id, minutes in minutes_map.items():
        emp = Employee.objects.get(pk=emp_id)

        line = existing.get(emp_id) or PayrollLine(
            period=period,
            employee=emp,
            adjustment=Decimal("0.00"),
        )

        hourly = _to_decimal(emp.hourly_rate, "0")
        base_amount = (Decimal(minutes) / Decimal(60)) * hourly

        line.minutes = int(minutes)
        line.hourly_rate = hourly
        line.currency = emp.currency or "USD"
        line.base_amount = base_amount.quantize(TWOPL, rounding=ROUND_HALF_UP)
        line.save()

    messages.success(request, "Nómina recalculada correctamente.")
    return redirect("payroll_detail", pk=period.pk)


@login_required
@user_passes_test(is_admin)
def payroll_detail(request, pk):
    period = get_object_or_404(PayrollPeriod, pk=pk)
    lines = list(period.lines.select_related("employee").all())

    line_id = request.GET.get("line")
    form = None

    if line_id:
        line = get_object_or_404(PayrollLine, pk=line_id, period=period)

        if request.method == "POST" and period.is_open:
            form = PayrollLineAdjustmentForm(request.POST, instance=line)

            if form.is_valid():
                form.save()
                messages.success(request, "Ajuste guardado correctamente.")
                return redirect("payroll_detail", pk=period.pk)

            messages.error(request, "No se pudo guardar el ajuste.")

        else:
            form = PayrollLineAdjustmentForm(instance=line)

    totals_by_currency = defaultdict(lambda: {
        "base": Decimal("0.00"),
        "adjustment": Decimal("0.00"),
        "total": Decimal("0.00"),
    })

    lines_fmt = []

    for line in lines:
        currency = line.currency or line.employee.currency or "USD"
        base = Decimal(line.base_amount or 0)
        adjustment = Decimal(line.adjustment or 0)
        total = (base + adjustment).quantize(TWOPL, rounding=ROUND_HALF_UP)

        totals_by_currency[currency]["base"] += base
        totals_by_currency[currency]["adjustment"] += adjustment
        totals_by_currency[currency]["total"] += total

        lines_fmt.append({
            "obj": line,
            "employee": line.employee,
            "currency": currency,
            "hours_hhmm": format_hhmm_from_minutes(line.minutes),
            "rate": format_money_latam(line.hourly_rate, currency),
            "base": format_money_latam(base, currency),
            "adj": format_money_latam(adjustment, currency),
            "total": format_money_latam_whole(total, currency),
            "total_cent": format_money_latam(total, currency),
        })

    employees = [line.employee for line in lines]

    att_map = defaultdict(int)

    if employees:
        att_qs = Attendance.objects.filter(
            employee_id__in=[e.id for e in employees],
            date__gte=period.start_date,
            date__lte=period.end_date,
        )

        if period.project_id:
            att_qs = att_qs.filter(project=period.project)

        for att in att_qs:
            att_map[(att.employee_id, att.date)] += _billable_minutes(att)

    weeks = []
    total_period_minutes = 0
    total_period_pay_by_currency = defaultdict(lambda: Decimal("0.00"))

    for w_start, w_end in _week_chunks(period.start_date, period.end_date):
        wk = _build_week_table(
            employees,
            att_map,
            w_start,
            w_end,
            period.start_date,
            period.end_date,
        )

        weeks.append(wk)

        total_period_minutes += int(wk["total_hours_all_raw"] * Decimal(60))

        for currency, value in wk["total_payroll_raw"].items():
            total_period_pay_by_currency[currency] += value

    return render(request, "attendance/payroll_detail.html", {
        "period": period,
        "lines": lines,
        "lines_fmt": lines_fmt,
        "form": form,
        "weeks": weeks,
        "title": f"{period.start_date} → {period.end_date}",
        "totals_by_currency": {
            currency: {
                "base": format_money_latam_whole(values["base"], currency),
                "adjustment": format_money_latam_whole(values["adjustment"], currency),
                "total": format_money_latam_whole(values["total"], currency),
            }
            for currency, values in totals_by_currency.items()
        },
        "grand_total_hours": format_hhmm_from_minutes(total_period_minutes),
        "grand_total_payroll": [
            format_money_latam_whole(value, currency)
            for currency, value in total_period_pay_by_currency.items()
        ],
    })


@login_required
@user_passes_test(is_admin)
def payroll_close(request, pk):
    period = get_object_or_404(PayrollPeriod, pk=pk)

    if request.method == "POST":
        if period.is_open:
            period.status = "CLOSED"
            period.closed_at = timezone.now()
            period.closed_by = request.user
            period.save(update_fields=["status", "closed_at", "closed_by"])
            messages.success(request, "Periodo cerrado correctamente.")
        else:
            messages.warning(request, "Este periodo ya estaba cerrado.")

    return redirect("payroll_detail", pk=period.pk)


# ──────────────────────────────────────────────────────────────────────────────
# Export Excel
# ──────────────────────────────────────────────────────────────────────────────

@login_required
@user_passes_test(is_admin)
def payroll_export_xlsx(request):
    form = PayrollExportForm(request.GET or None)
    period_pk = request.GET.get("period_pk")

    if not period_pk and not (form.is_bound and form.is_valid()):
        return render(request, "attendance/payroll_export_filter.html", {"form": form})

    wb = Workbook()
    ws = wb.active
    ws.title = "Nomina"

    header_font = Font(bold=True)
    center = Alignment(horizontal="center")
    right = Alignment(horizontal="right")

    def append_header(sheet, values):
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.font = header_font
            cell.alignment = center

    def auto_width(sheet):
        for col_cells in sheet.columns:
            length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
            sheet.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 50)

    if period_pk:
        period = get_object_or_404(PayrollPeriod, pk=period_pk)
        lines = list(period.lines.select_related("employee").order_by("employee__full_name"))

        ws.append([f"Nómina {period.start_date} → {period.end_date}"])
        ws.append([f"Proyecto: {period.project.name if period.project else 'Todos'}"])
        ws.append([])

        append_header(ws, ["Documento", "Empleado", "Horas HH:MM", "Horas dec.", "Tarifa", "Moneda", "Base", "Ajuste", "Total"])

        totals_by_currency = defaultdict(lambda: Decimal("0.00"))

        for line in lines:
            currency = line.currency or line.employee.currency or "USD"
            hours_dec = Decimal(line.minutes or 0) / Decimal(60)
            total = Decimal(line.base_amount or 0) + Decimal(line.adjustment or 0)

            ws.append([
                line.employee.document_id or "",
                line.employee.full_name,
                format_hhmm_from_minutes(line.minutes),
                float(hours_dec),
                float(line.hourly_rate or 0),
                currency,
                float(line.base_amount or 0),
                float(line.adjustment or 0),
                float(total),
            ])

            for col in [4, 5, 7, 8, 9]:
                ws.cell(ws.max_row, col).alignment = right

            totals_by_currency[currency] += total

        ws.append([])
        ws.append(["TOTALES"])
        ws.cell(ws.max_row, 1).font = header_font

        for currency, value in totals_by_currency.items():
            ws.append(["", "", "", "", "", currency, "", "", float(value)])

        auto_width(ws)

        filename = f"nomina_periodo_{period.start_date}_{period.end_date}.xlsx"

    else:
        project = form.cleaned_data.get("project")
        scope = form.cleaned_data["scope"]
        month = form.cleaned_data.get("month")
        year = form.cleaned_data.get("year")
        half = form.cleaned_data.get("half")
        today = timezone.localdate()

        if scope == "mes":
            y = month.year if month else today.year
            m = month.month if month else today.month
            start, end = date(y, m, 1), date(y, m, monthrange(y, m)[1])
            title = f"Nómina {y}-{m:02d}"

        elif scope == "anio":
            y = year or today.year
            start, end = date(y, 1, 1), date(y, 12, 31)
            title = f"Nómina {y}"

        else:
            y = month.year if month else today.year
            m = month.month if month else today.month
            last_day = monthrange(y, m)[1]

            if half == "1" or (not half and today.day <= 15):
                start, end = date(y, m, 1), date(y, m, 15)
                half = "1"
            else:
                start, end = date(y, m, 16), date(y, m, last_day)
                half = "2"

            title = f"Nómina {y}-{m:02d} Q{half}"

        qs = Attendance.objects.select_related("employee", "project").filter(date__range=(start, end))

        if project:
            qs = qs.filter(project=project)

        agg = defaultdict(lambda: {
            "document": "",
            "name": "",
            "minutes": 0,
            "rate": Decimal("0.00"),
            "currency": "USD",
        })

        for att in qs:
            key = att.employee_id
            agg[key]["document"] = att.employee.document_id or ""
            agg[key]["name"] = att.employee.full_name
            agg[key]["minutes"] += _billable_minutes(att)
            agg[key]["rate"] = att.employee.hourly_rate or Decimal("0.00")
            agg[key]["currency"] = att.employee.currency or "USD"

        ws.append([title])
        ws.append([f"Rango: {start} → {end} | Proyecto: {project.name if project else 'Todos'}"])
        ws.append([])

        append_header(ws, ["Documento", "Empleado", "Horas HH:MM", "Horas dec.", "Tarifa", "Moneda", "Total"])

        totals_by_currency = defaultdict(lambda: Decimal("0.00"))

        for _, data in sorted(agg.items(), key=lambda x: x[1]["name"].lower()):
            hours_dec = Decimal(data["minutes"]) / Decimal(60)
            total = (hours_dec * data["rate"]).quantize(TWOPL, rounding=ROUND_HALF_UP)
            currency = data["currency"]

            ws.append([
                data["document"],
                data["name"],
                format_hhmm_from_minutes(data["minutes"]),
                float(hours_dec),
                float(data["rate"]),
                currency,
                float(total),
            ])

            totals_by_currency[currency] += total

        ws.append([])
        ws.append(["TOTALES"])
        ws.cell(ws.max_row, 1).font = header_font

        for currency, value in totals_by_currency.items():
            ws.append(["", "", "", "", "", currency, float(value)])

        auto_width(ws)
        filename = f"nomina_{start}_{end}_{getattr(project, 'slug', 'todos')}.xlsx"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)

    return response


# ──────────────────────────────────────────────────────────────────────────────
# Finanzas de proyecto
# ──────────────────────────────────────────────────────────────────────────────

@login_required
@user_passes_test(is_admin)
def project_new(request):
    if request.method == "POST":
        form = ProjectForm(request.POST)

        if form.is_valid():
            project = form.save()
            messages.success(request, "Proyecto creado correctamente.")
            return redirect("project_finance_dashboard", project_slug=project.slug)

        messages.error(request, "No se pudo crear el proyecto. Revisa los campos.")

    else:
        form = ProjectForm()

    return render(request, "attendance/project_new.html", {"form": form})


@login_required
@user_passes_test(is_admin)
def income_new(request):
    if request.method == "POST":
        form = IncomeForm(request.POST)

        if form.is_valid():
            income = form.save(commit=False)
            income.created_by = request.user
            income.save()
            messages.success(request, "Ingreso registrado correctamente.")
            return redirect("project_finance_dashboard", project_slug=income.project.slug)

        messages.error(request, "No se pudo registrar el ingreso.")

    else:
        form = IncomeForm()

    return render(request, "attendance/income_form.html", {"form": form})


@login_required
@user_passes_test(is_admin)
def expense_new(request):
    if request.method == "POST":
        form = ExpenseForm(request.POST)

        if form.is_valid():
            expense = form.save(commit=False)
            expense.created_by = request.user
            expense.save()
            messages.success(request, "Gasto registrado correctamente.")
            return redirect("project_finance_dashboard", project_slug=expense.project.slug)

        messages.error(request, "No se pudo registrar el gasto.")

    else:
        form = ExpenseForm()

    return render(request, "attendance/expense_form.html", {"form": form})


@login_required
@user_passes_test(is_admin)
def project_finance_dashboard(request, project_slug):

    project = get_object_or_404(
        Project,
        slug=project_slug
    )

    # =====================================================
    # FILTROS
    # =====================================================

    month = (request.GET.get("month") or "").strip()
    year = (request.GET.get("year") or "").strip()

    incomes = ProjectIncome.objects.filter(project=project)
    expenses = ProjectExpense.objects.filter(project=project)

    # =====================================================
    # FILTRO POR MES
    # =====================================================

    if month:

        try:
            year_part, month_part = month.split("-")

            incomes = incomes.filter(
                date__year=int(year_part),
                date__month=int(month_part)
            )

            expenses = expenses.filter(
                date__year=int(year_part),
                date__month=int(month_part)
            )

        except Exception:
            messages.warning(
                request,
                "El formato del mes no es válido."
            )

    # =====================================================
    # FILTRO SOLO POR AÑO
    # =====================================================

    elif year:

        try:

            year_int = int(year)

            incomes = incomes.filter(
                date__year=year_int
            )

            expenses = expenses.filter(
                date__year=year_int
            )

        except ValueError:

            messages.warning(
                request,
                "El año ingresado no es válido."
            )

    # =====================================================
    # ORDEN
    # =====================================================

    incomes = incomes.order_by("-date", "-id")
    expenses = expenses.order_by("-date", "-id")

    # =====================================================
    # TOTALES
    # =====================================================

    inc_total = (
        incomes.aggregate(
            total=Sum("amount")
        )["total"]
        or 0
    )

    exp_total = (
        expenses.aggregate(
            total=Sum("amount")
        )["total"]
        or 0
    )

# =====================================================
# COSTO NÓMINA
# =====================================================

    attendance_qs = Attendance.objects.select_related(
        "employee"
    ).filter(
        project=project
    )

    if month:

        try:

            attendance_qs = attendance_qs.filter(
                date__year=int(year_part),
                date__month=int(month_part)
            )

        except Exception:
            pass

    elif year:

        try:

            attendance_qs = attendance_qs.filter(
                date__year=int(year)
            )

        except Exception:
            pass

    payroll_cost = Decimal("0.00")

    for att in attendance_qs:

        minutes = Decimal(_billable_minutes(att))

        rate = att.employee.hourly_rate or Decimal("0.00")

        payroll_cost += (
            (minutes / Decimal("60")) * rate
        )

    payroll_cost = payroll_cost.quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP
    )
    # =====================================================
    # BALANCE
    # =====================================================

    contract_value = project.contract_value or 0

    balance = (
        contract_value
        + inc_total
        - exp_total
        - payroll_cost
    )

    # =====================================================
    # GRÁFICAS
    # =====================================================

    line_values = [
        float(contract_value),
        float(contract_value + inc_total),
        float(contract_value + inc_total - exp_total),
        float(contract_value + inc_total - exp_total - payroll_cost),
        float(balance),
    ]

    chart_values = [
        float(contract_value),
        float(inc_total),
        float(exp_total),
        float(payroll_cost),
        float(balance),
    ]

    # =====================================================
    # FORMATO
    # =====================================================

    currency = (
        project.contract_currency
        or "USD"
    )

    def fmt_money(value):

        try:

            return (
                f"{currency} "
                f"{float(value):,.2f}"
            )

        except Exception:

            return f"{currency} 0.00"

    # =====================================================
    # CONTEXTO
    # =====================================================

    context = {

        "project": project,

        "currency": currency,

        "month": month,
        "year": year,

        "incomes": incomes[:15],
        "expenses": expenses[:15],

        "contract_value": contract_value,
        "inc_total": inc_total,
        "exp_total": exp_total,
        "payroll_cost": payroll_cost,
        "balance": balance,

        "contract_value_fmt": fmt_money(contract_value),
        "inc_total_fmt": fmt_money(inc_total),
        "exp_total_fmt": fmt_money(exp_total),
        "payroll_cost_fmt": fmt_money(payroll_cost),
        "balance_fmt": fmt_money(balance),

        "line_values": line_values,
        "chart_values": chart_values,
    }

    return render(
        request,
        "attendance/project_finance_dashboard.html",
        context
    )
# ──────────────────────────────────────────────────────────────────────────────
# Export PDF simple
# ──────────────────────────────────────────────────────────────────────────────

@login_required
@user_passes_test(is_admin)
def payroll_export_pdf(request):
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    from reportlab.lib.units import cm

    form = PayrollExportForm(request.GET or None)
    period_pk = request.GET.get("period_pk")

    if not period_pk and not (form.is_bound and form.is_valid()):
        return render(request, "attendance/payroll_export_filter.html", {"form": form})

    buf = BytesIO()

    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
    )

    styles = getSampleStyleSheet()

    h1 = ParagraphStyle(
        "H1",
        parent=styles["Heading1"],
        fontSize=14,
        textColor=colors.HexColor("#dc3545"),
        spaceAfter=8,
    )

    small = ParagraphStyle(
        "SMALL",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.gray,
    )

    story = []

    def table_style():
        return TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f8d7da")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#dc3545")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
            ("ALIGN", (0, 1), (1, -1), "LEFT"),
        ])

    if period_pk:
        period = get_object_or_404(PayrollPeriod, pk=period_pk)
        lines = period.lines.select_related("employee").order_by("employee__full_name")

        story.append(Paragraph(f"Nómina {period.start_date} → {period.end_date}", h1))
        story.append(Paragraph(f"Proyecto: {period.project.name if period.project else 'Todos'}", small))
        story.append(Spacer(1, 0.3 * cm))

        data = [["Documento", "Empleado", "Horas", "Tarifa", "Moneda", "Base", "Ajuste", "Total"]]

        for line in lines:
            currency = line.currency or line.employee.currency or "USD"
            total = Decimal(line.base_amount or 0) + Decimal(line.adjustment or 0)

            data.append([
                line.employee.document_id or "",
                line.employee.full_name,
                format_hhmm_from_minutes(line.minutes),
                format_money_latam(line.hourly_rate, currency),
                currency,
                format_money_latam(line.base_amount, currency),
                format_money_latam(line.adjustment, currency),
                format_money_latam_whole(total, currency),
            ])

        filename = f"nomina_periodo_{period.start_date}_{period.end_date}.pdf"

    else:
        return redirect("payroll_export_xlsx")

    t = Table(data, repeatRows=1, hAlign="CENTER")
    t.setStyle(table_style())
    story.append(t)

    doc.build(story)

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.write(buf.getvalue())
    buf.close()

    return response